#!/usr/bin/env python3
"""paperless-export — export a paperless-ngx document list to a sheet you can check.

Read-only against the instance. Python 3.10+. Stdlib only, except openpyxl for --format xlsx.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import html
import ipaddress
import json
import os
import re
import socket
import sys
import time
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass, field as dc_field
from pathlib import Path
from typing import Any

__version__ = "0.4.0"

API_VERSION = "10"
PAGE_SIZE = 1000
TIMEOUT = 60

EXIT_OK = 0
EXIT_FAIL = 1
EXIT_USAGE = 2
EXIT_RECONCILE = 3


class ExportError(Exception):
    """Fatal, expected failure. Message is shown to the user; exit 1."""


class UsageError(Exception):
    """Bad invocation; exit 2."""


class ReconcileError(Exception):
    """Row count disagreed with the API count; exit 3."""


# --------------------------------------------------------------------------------------
# saved view filter rules: rule_type -> (query param, isnull param or None, joins multiple
# values with commas). The one source; the JavaScript engine's copy is generated from this.
# --------------------------------------------------------------------------------------
RULE_MAP: dict[int, tuple[str, str | None, bool]] = {
    0: ("title__icontains", None, False),
    1: ("content__icontains", None, False),
    2: ("archive_serial_number", None, False),
    3: ("correspondent__id", "correspondent__isnull", False),
    4: ("document_type__id", "document_type__isnull", False),
    5: ("is_in_inbox", None, False),
    6: ("tags__id__all", None, True),
    7: ("is_tagged", None, False),
    8: ("created__date__lt", None, False),
    9: ("created__date__gt", None, False),
    10: ("created__year", None, False),
    11: ("created__month", None, False),
    12: ("created__day", None, False),
    13: ("added__date__lt", None, False),
    14: ("added__date__gt", None, False),
    15: ("modified__date__lt", None, False),
    16: ("modified__date__gt", None, False),
    17: ("tags__id__none", None, True),
    18: ("archive_serial_number__isnull", None, False),
    19: ("title_content", None, False),
    20: ("query", None, False),
    21: ("more_like_id", None, False),
    22: ("tags__id__in", None, True),
    23: ("archive_serial_number__gt", None, False),
    24: ("archive_serial_number__lt", None, False),
    25: ("storage_path__id", "storage_path__isnull", False),
    26: ("correspondent__id__in", None, True),
    27: ("correspondent__id__none", None, True),
    28: ("document_type__id__in", None, True),
    29: ("document_type__id__none", None, True),
    30: ("storage_path__id__in", None, True),
    31: ("storage_path__id__none", None, True),
    32: ("owner__id", None, False),
    33: ("owner__id__in", None, True),
    34: ("owner__isnull", None, False),
    35: ("owner__id__none", None, True),
    36: ("custom_fields__icontains", None, False),
    37: ("shared_by__id", None, True),
    38: ("custom_fields__id__all", None, True),
    39: ("custom_fields__id__in", None, True),
    40: ("custom_fields__id__none", None, True),
    41: ("has_custom_fields", None, False),
    42: ("custom_field_query", None, False),
    43: ("created__date__lte", None, False),
    44: ("created__date__gte", None, False),
    45: ("added__date__lte", None, False),
    46: ("added__date__gte", None, False),
    47: ("mime_type", None, False),
    48: ("title_search", None, False),
    49: ("text", None, False),
    50: ("has_duplicates", None, False),
}

# saved view display_fields -> our column key
VIEW_FIELD_MAP = {
    "title": "title",
    "created": "created",
    "added": "added",
    "modified": "modified",
    "tag": "tags",
    "correspondent": "correspondent",
    "documenttype": "document_type",
    "storagepath": "storage_path",
    "note": "notes",
    "owner": "owner",
    "shared": "shared",
    "asn": "archive_serial_number",
    "pagecount": "page_count",
}

DEFAULT_COLUMNS = [
    "id", "title", "archive_serial_number", "correspondent", "document_type", "tags",
    "storage_path", "created", "added", "modified", "original_file_name",
    "archived_file_name", "page_count", "mime_type", "notes", "versions",
]

SUMMABLE = {"integer", "float", "monetary"}
INJECTION_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

# --date presets. Anything else is a strftime pattern, validated before the export runs.
DATE_PRESETS = {"iso": "%Y-%m-%d", "us": "%m/%d/%Y", "eu": "%d.%m.%Y"}
DATE_PROBE = dt.date(2026, 3, 4)  # day and month differ, so a swapped pattern is visible


# --------------------------------------------------------------------------------------
# config
# --------------------------------------------------------------------------------------
@dataclass
class Config:
    base_url: str
    token: str
    view: str | None = None
    filters: dict[str, Any] = dc_field(default_factory=dict)
    sums: list[str] = dc_field(default_factory=list)
    columns: list[str] | None = None
    content: bool = False
    separator: str = ","
    decimal: str = "."
    date_spec: str = "iso"          # as the user typed it; the trailer records this
    date_fmt: str = "%Y-%m-%d"      # resolved and validated strftime pattern
    headers: str | None = None      # None: readable for csv/xlsx, raw for json/jsonl
    list_sep: str = "; "
    output: Path = Path("export")
    formats: list[str] = dc_field(default_factory=lambda: ["csv"])
    pack: bool = False
    pack_version: str = "latest"
    inventory: bool = False
    insecure: bool = False


# --------------------------------------------------------------------------------------
# transport
# --------------------------------------------------------------------------------------
def _resolve_host(host: str) -> list[str]:
    """Every address a name resolves to, both families. Empty list means it does not resolve."""
    try:
        infos = socket.getaddrinfo(host, None, proto=socket.IPPROTO_TCP)
    except (OSError, UnicodeError):
        return []
    return [info[4][0] for info in infos]


def _addr_is_private(addr: str) -> bool:
    addr = addr.split("%", 1)[0]  # drop an IPv6 scope id
    try:
        ip = ipaddress.ip_address(addr)
    except ValueError:
        return False
    mapped = getattr(ip, "ipv4_mapped", None)
    if mapped is not None:
        ip = mapped
    return ip.is_loopback or ip.is_private or ip.is_link_local


def is_private_host(host: str) -> bool:
    """A bare LAN name is as private as the address behind it. Resolve, then judge.

    Private only when *every* address the name resolves to is loopback, RFC1918, ULA or
    link-local; a name that does not resolve is not private (never assume in our favour).
    """
    if not host:
        return False
    host = host.lower().rstrip(".")
    if host == "localhost" or host.endswith(".localhost") or host.endswith(".local"):
        return True
    if _addr_is_private(host):
        return True
    try:
        ipaddress.ip_address(host.split("%", 1)[0])
    except ValueError:
        pass
    else:
        return False  # a literal address that is not private; do not resolve it
    addrs = _resolve_host(host)
    return bool(addrs) and all(_addr_is_private(a) for a in addrs)


def origin_of(url: str) -> tuple[str, str, int | None]:
    p = urllib.parse.urlsplit(url)
    port = p.port
    if port is None:
        port = {"http": 80, "https": 443}.get(p.scheme)
    return (p.scheme.lower(), (p.hostname or "").lower(), port)


class SameOriginRedirect(urllib.request.HTTPRedirectHandler):
    """Refuse any redirect that leaves the origin the user typed.

    urllib copies every header except content-length/content-type onto the redirected
    request, Authorization included, so an instance that is compromised or MITM'd on a
    plain-http LAN could hand our token to any host it names. Same-origin only, always.
    """

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        here, there = origin_of(req.full_url), origin_of(newurl)
        if here != there:
            if here[0] == "https" and there[0] == "http":
                raise ExportError(f"refused https->http redirect to {newurl}")
            raise ExportError(
                f"refused a redirect off the instance you named: {req.full_url} -> {newurl}. "
                "The token travels with a redirect; it does not leave your instance."
            )
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def check_url(raw: str, insecure: bool) -> str:
    parts = urllib.parse.urlsplit(raw)
    if parts.scheme not in ("http", "https"):
        raise UsageError(f"URL scheme must be http or https, got {parts.scheme!r}")
    if not parts.netloc:
        raise UsageError(f"URL has no host: {raw!r}")
    if parts.query or parts.fragment:
        raise UsageError("refusing a URL carrying a query or fragment (tokens never ride in URLs)")
    if parts.scheme == "http" and not is_private_host(parts.hostname or ""):
        if not insecure:
            raise UsageError(
                f"refusing plain http to {parts.hostname!r}: it is not a private address "
                "and does not resolve to one. Use https, or pass --insecure if you "
                "accept a cleartext token"
            )
        warn(f"plain http to public host {parts.hostname} (--insecure)")
    elif parts.scheme == "http":
        warn(f"plain http to private host {parts.hostname}")
    return f"{parts.scheme}://{parts.netloc}{parts.path.rstrip('/')}"


class Client:
    def __init__(self, base: str, token: str):
        self.base = base
        self.token = token
        self.opener = urllib.request.build_opener(SameOriginRedirect)

    def get(self, path: str, params: dict[str, Any] | None = None) -> tuple[Any, Any]:
        url = f"{self.base}{path}"
        if params:
            url += ("&" if "?" in path else "?") + urllib.parse.urlencode(params, doseq=False)
        req = urllib.request.Request(
            url,
            headers={
                "Authorization": f"Token {self.token}",
                "Accept": f"application/json; version={API_VERSION}",
            },
        )
        last: Exception | None = None
        for attempt in (1, 2):
            try:
                with self.opener.open(req, timeout=TIMEOUT) as r:
                    return json.loads(r.read()), r.headers
            except urllib.error.HTTPError as e:
                detail = e.read().decode(errors="replace")[:300]
                # 4xx is the caller's fault; retrying cannot help.
                raise ExportError(f"{path} returned HTTP {e.code}: {detail}") from None
            except urllib.error.URLError as e:
                last = ExportError(f"{path} unreachable: {e.reason}")
                if attempt == 1:
                    time.sleep(1)
        raise last  # type: ignore[misc]

    def get_bytes(self, path: str, params: dict[str, Any] | None = None) -> bytes:
        url = f"{self.base}{path}"
        if params:
            url += "?" + urllib.parse.urlencode(params)
        req = urllib.request.Request(url, headers={"Authorization": f"Token {self.token}"})
        try:
            with self.opener.open(req, timeout=TIMEOUT) as r:
                return r.read()
        except urllib.error.HTTPError as e:
            raise ExportError(f"{path} returned HTTP {e.code}") from None
        except urllib.error.URLError as e:
            raise ExportError(f"{path} unreachable: {e.reason}") from None

    def next_path(self, nxt: str) -> str:
        """Turn a `next` link into a path on *our* base, or refuse it.

        paperless builds `next` from its own configured hostname, which is often not the
        one the user typed. Take the path and query; refuse an origin we did not choose,
        so a hostile or misconfigured response cannot walk the token to another host.
        """
        p = urllib.parse.urlsplit(nxt)
        if p.scheme or p.netloc:
            if origin_of(nxt) != origin_of(self.base):
                raise ExportError(
                    f"the API's next-page link points at another origin ({p.scheme}://{p.netloc}); "
                    "refusing to follow it. Check PAPERLESS_URL on the server."
                )
        return urllib.parse.urlunsplit(("", "", p.path, p.query, ""))

    def get_all(self, path: str) -> list[dict]:
        """Page through a list endpoint, returning every result."""
        out: list[dict] = []
        data, _ = self.get(path, {"page_size": PAGE_SIZE})
        out.extend(data.get("results", []))
        nxt = data.get("next")
        while nxt:
            data, _ = self.get(self.next_path(nxt))
            out.extend(data.get("results", []))
            nxt = data.get("next")
        return out


MIN_SECRET_LEN = 8
_SECRETS: list[str] = []


def redact(msg: str) -> str:
    """Belt and braces: no string we ever print may contain the token.

    Nothing is *supposed* to put it there, but a server can echo a header back in an error
    body and we print 300 characters of that body. One place to enforce it, so it holds.
    """
    # Only substantial strings: a real paperless token is 40 hex characters, and blanking
    # a two-character "token" would redact the whole message into nonsense.
    for secret in _SECRETS:
        if len(secret) >= MIN_SECRET_LEN and secret in msg:
            msg = msg.replace(secret, "***REDACTED***")
    return msg


def warn(msg: str) -> None:
    print(f"note: {redact(msg)}", file=sys.stderr)


# --------------------------------------------------------------------------------------
# lookups
# --------------------------------------------------------------------------------------
class Lookups:
    def __init__(self, client: Client):
        self.correspondents = {r["id"]: r["name"] for r in client.get_all("/api/correspondents/")}
        self.document_types = {r["id"]: r["name"] for r in client.get_all("/api/document_types/")}
        self.tags = {r["id"]: r["name"] for r in client.get_all("/api/tags/")}
        self.storage_paths = {r["id"]: r["name"] for r in client.get_all("/api/storage_paths/")}
        # Served alphabetically by name (viewset order_by("name")); that is the column order.
        self.custom_fields = client.get_all("/api/custom_fields/")
        self.cf_by_id = {f["id"]: f for f in self.custom_fields}

    def resolve(self, table: dict[int, str], value: str, what: str) -> int:
        if str(value).isdigit():
            vid = int(value)
            if vid in table:
                return vid
            raise UsageError(f"no {what} with id {vid}")
        matches = [i for i, n in table.items() if n.lower() == str(value).lower()]
        if len(matches) == 1:
            return matches[0]
        if not matches:
            near = [n for n in table.values() if str(value).lower() in n.lower()][:5]
            hint = f" did you mean: {', '.join(near)}?" if near else ""
            raise UsageError(f"no {what} named {value!r}.{hint}")
        raise UsageError(f"{what} {value!r} is ambiguous (ids {matches})")

    def cf_column_name(self, f: dict) -> str:
        """Two fields may share a name; disambiguate with the id when they do."""
        same = [g for g in self.custom_fields if g["name"] == f["name"]]
        return f"custom_field:{f['name']}" if len(same) == 1 else f"custom_field:{f['name']}#{f['id']}"


# --------------------------------------------------------------------------------------
# filter resolution
# --------------------------------------------------------------------------------------
def params_from_view(view: dict, notes: list[str]) -> tuple[dict[str, str], str | None, list[str] | None]:
    multi: dict[str, list[str]] = {}
    params: dict[str, str] = {}
    for rule in view.get("filter_rules") or []:
        rt = rule.get("rule_type")
        val = rule.get("value")
        if rt not in RULE_MAP:
            raise ExportError(
                f"saved view uses unknown filter rule_type {rt}. This paperless version added a "
                "rule this tool does not map. Refusing rather than silently dropping a filter "
                "(that would produce a sheet with too many rows that still reconciles). "
                f"Report rule type {rt} to the maintainer so it can be added."
            )
        param, isnull_param, is_multi = RULE_MAP[rt]
        if val is None and isnull_param:
            params[isnull_param] = "true"
            continue
        if val is None:
            continue
        if is_multi:
            multi.setdefault(param, []).append(str(val))
        else:
            params[param] = str(val)
    for param, vals in multi.items():
        params[param] = ",".join(vals)

    sort = None
    if view.get("sort_field"):
        sort = ("-" if view.get("sort_reverse") else "") + str(view["sort_field"])

    columns = None
    display = view.get("display_fields") or []
    if display:
        columns = []
        for d in display:
            if d in VIEW_FIELD_MAP:
                columns.append(VIEW_FIELD_MAP[d])
            elif d.startswith("custom_field_"):
                columns.append(d)  # resolved to a name later, once lookups are known
            else:
                notes.append(f"saved view display field {d!r} not recognised, column skipped")
    return params, sort, columns


def params_from_flags(cfg: Config, lk: Lookups) -> dict[str, str]:
    f = cfg.filters
    params: dict[str, str] = {}
    if f.get("query"):
        params["query"] = f["query"]
    if f.get("title"):
        params["title__icontains"] = f["title"]
    for key, table, param in (
        ("correspondent", lk.correspondents, "correspondent__id__in"),
        ("document_type", lk.document_types, "document_type__id__in"),
        ("storage_path", lk.storage_paths, "storage_path__id__in"),
    ):
        if f.get(key):
            ids = [str(lk.resolve(table, v, key)) for v in f[key]]
            params[param] = ",".join(ids)
    if f.get("tags"):
        ids = [str(lk.resolve(lk.tags, v, "tag")) for v in f["tags"]]
        params["tags__id__all"] = ",".join(ids)
    if f.get("created_from"):
        params["created__date__gte"] = f["created_from"]
    if f.get("created_to"):
        params["created__date__lte"] = f["created_to"]
    if f.get("custom_field_query"):
        params["custom_field_query"] = f["custom_field_query"]
    return params


# --------------------------------------------------------------------------------------
# row building
# --------------------------------------------------------------------------------------
def parse_monetary(raw: Any) -> tuple[str, float | None]:
    """'USD123.45' -> ('USD', 123.45). Currency may be absent."""
    if raw in (None, ""):
        return "", None
    s = str(raw)
    m = re.match(r"^([A-Za-z]{3})?\s*(-?[\d.,]+)$", s.strip())
    if not m:
        return "", None
    cur = (m.group(1) or "").upper()
    num = m.group(2).replace(",", "")
    try:
        return cur, float(num)
    except ValueError:
        return cur, None


def resolve_date_format(spec: str) -> str:
    """Presets, or a strftime pattern proved on a fixed date before a single row is fetched."""
    fmt = DATE_PRESETS.get(spec, spec)
    try:
        out = DATE_PROBE.strftime(fmt)
    except (ValueError, TypeError) as e:
        raise UsageError(f"--date: {spec!r} is not a usable strftime pattern ({e})") from None
    if not out:
        raise UsageError(f"--date: {spec!r} produces an empty date")
    if not any(ch.isdigit() for ch in out):
        # macOS strftime swallows an unknown specifier ('%Q' -> 'Q'), so a typo would quietly
        # stamp the same nonsense into every date column. A date has a number in it.
        raise UsageError(
            f"--date: {spec!r} produces {out!r}, which has no number in it. That is not a date."
        )
    if out.startswith(INJECTION_PREFIXES):
        raise UsageError(
            f"--date: {spec!r} produces {out!r}, which a spreadsheet reads as a formula. "
            "Presentation is adjustable; the formula guard is not."
        )
    return fmt


def fmt_date(value: Any, fmt: str) -> str:
    if not value:
        return ""
    s = str(value)
    try:
        d = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        return s[:10]
    return d.strftime(fmt)


def flatten_custom_fields(doc: dict, lk: Lookups, cfg: Config, notes: list[str]) -> dict[str, Any]:
    """One column per defined field, alphabetical by name; monetary also gets :currency."""
    out: dict[str, Any] = {}
    for f in lk.custom_fields:
        col = lk.cf_column_name(f)
        out[col] = ""
        if f["data_type"] == "monetary":
            out[f"{col}:currency"] = ""

    values = {v["field"]: v.get("value") for v in (doc.get("custom_fields") or [])}
    for f in lk.custom_fields:
        if f["id"] not in values:
            continue
        col = lk.cf_column_name(f)
        val = values[f["id"]]
        dtype = f["data_type"]
        if val is None:
            continue
        if dtype in ("string", "longtext", "url"):
            out[col] = val
        elif dtype == "integer":
            out[col] = val
        elif dtype == "float":
            out[col] = val
        elif dtype == "boolean":
            out[col] = val
        elif dtype == "date":
            out[col] = fmt_date(val, cfg.date_fmt)
        elif dtype == "monetary":
            cur, num = parse_monetary(val)
            out[col] = num if num is not None else val
            out[f"{col}:currency"] = cur
        elif dtype == "select":
            opts = (f.get("extra_data") or {}).get("select_options") or []
            label = next((o.get("label") for o in opts if o.get("id") == val), None)
            if label is None:
                out[col] = val
                notes.append(f"select field {f['name']!r}: option id {val!r} has no label, written verbatim")
            else:
                out[col] = label
        elif dtype == "documentlink":
            out[col] = cfg.list_sep.join(str(x) for x in (val or []))
        else:
            out[col] = val
            notes.append(f"custom field {f['name']!r} has unknown data_type {dtype!r}, written verbatim")
    return out


def build_row(doc: dict, lk: Lookups, cfg: Config, notes: list[str]) -> dict[str, Any]:
    tags = cfg.list_sep.join(lk.tags.get(t, str(t)) for t in (doc.get("tags") or []))
    row: dict[str, Any] = {
        "id": doc.get("id"),
        "title": doc.get("title") or "",
        "archive_serial_number": doc.get("archive_serial_number") if doc.get("archive_serial_number") is not None else "",
        "correspondent": lk.correspondents.get(doc.get("correspondent"), "") if doc.get("correspondent") else "",
        "document_type": lk.document_types.get(doc.get("document_type"), "") if doc.get("document_type") else "",
        "tags": tags,
        "storage_path": lk.storage_paths.get(doc.get("storage_path"), "") if doc.get("storage_path") else "",
        "created": fmt_date(doc.get("created") or doc.get("created_date"), cfg.date_fmt),
        "added": fmt_date(doc.get("added"), cfg.date_fmt),
        "modified": fmt_date(doc.get("modified"), cfg.date_fmt),
        "original_file_name": doc.get("original_file_name") or "",
        "archived_file_name": doc.get("archived_file_name") or "",
        "page_count": doc.get("page_count") if doc.get("page_count") is not None else "",
        "mime_type": doc.get("mime_type") or "",
        "notes": len(doc.get("notes") or []),
        "versions": len(doc.get("versions") or []),
        "owner": doc.get("owner") if doc.get("owner") is not None else "",
        "shared": doc.get("is_shared_by_requester", ""),
        "file": "",
        "file:version": "",
        "url": f"{cfg.base_url}/documents/{doc.get('id')}/details",
    }
    if cfg.content:
        row["content"] = doc.get("content") or ""
    row.update(flatten_custom_fields(doc, lk, cfg, notes))
    return row


def resolve_columns(cfg: Config, lk: Lookups, view_columns: list[str] | None, notes: list[str]) -> list[str]:
    cf_cols: list[str] = []
    for f in lk.custom_fields:
        col = lk.cf_column_name(f)
        cf_cols.append(col)
        if f["data_type"] == "monetary":
            cf_cols.append(f"{col}:currency")

    if view_columns is not None:
        cols = []
        for c in view_columns:
            if c.startswith("custom_field_"):
                try:
                    fid = int(c.rsplit("_", 1)[1])
                except ValueError:
                    notes.append(f"saved view column {c!r} unparseable, skipped")
                    continue
                f = lk.cf_by_id.get(fid)
                if not f:
                    notes.append(f"saved view references custom field id {fid}, which no longer exists; column skipped")
                    continue
                cols.append(lk.cf_column_name(f))
                if f["data_type"] == "monetary":
                    cols.append(f"{lk.cf_column_name(f)}:currency")
            else:
                cols.append(c)
        return cols

    if cfg.columns:
        known = set(DEFAULT_COLUMNS) | set(cf_cols) | {"owner", "shared", "file", "file:version", "url", "content"}
        for c in cfg.columns:
            if c not in known:
                raise UsageError(f"unknown column {c!r}")
        return list(cfg.columns)

    cols = list(DEFAULT_COLUMNS) + cf_cols
    if cfg.content:
        cols.append("content")
    if cfg.pack:
        cols += ["file", "file:version"]
    cols.append("url")
    return cols


# --------------------------------------------------------------------------------------
# fetch
# --------------------------------------------------------------------------------------
def fetch_documents(client: Client, params: dict[str, str], cfg: Config) -> tuple[list[dict], int]:
    q = dict(params)
    q["page_size"] = str(PAGE_SIZE)
    if not cfg.content:
        q["truncate_content"] = "true"

    data, _ = client.get("/api/documents/", q)
    count = data.get("count", 0)
    docs = list(data.get("results", []))
    nxt = data.get("next")
    while nxt:
        page, _ = client.get(client.next_path(nxt))
        if page.get("count", count) != count:
            raise ExportError(
                f"the document count changed mid-export ({count} -> {page.get('count')}). "
                "Something is writing to the instance. Re-run when it is quiet."
            )
        docs.extend(page.get("results", []))
        nxt = page.get("next")
    return docs, count


# --------------------------------------------------------------------------------------
# writers
# --------------------------------------------------------------------------------------
def readable_header(col: str) -> str:
    """Raw column name -> the heading a person reads. A rule, not a lookup table."""
    if col.startswith("custom_field:"):
        # A paperless field name is the user's own text; it is shown exactly as they typed it.
        rest = col[len("custom_field:"):]
        return rest[: -len(":currency")] + " currency" if rest.endswith(":currency") else rest
    s = col.replace(":", " ").replace("_", " ")
    return s[:1].upper() + s[1:]


def header_style(cfg: Config, fmt: str) -> str:
    """Unset means readable for the sheet formats, raw for the machine ones."""
    if cfg.headers:
        return cfg.headers
    return "readable" if fmt in ("csv", "xlsx") else "raw"


def readable_clashes(columns: list[str]) -> list[str]:
    """A custom field named 'Notes' reads the same as the built-in notes count. Raw names never do."""
    labels = [readable_header(c) for c in columns]
    return sorted({l for l in labels if labels.count(l) > 1})


def headings(columns: list[str], style: str) -> list[str]:
    if style != "readable":
        return list(columns)
    clash = set(readable_clashes(columns))
    return [c if readable_header(c) in clash else readable_header(c) for c in columns]


def is_formula_bait(value: Any) -> bool:
    return isinstance(value, str) and value.startswith(INJECTION_PREFIXES)


def guard(value: Any) -> Any:
    """CSV only: prefix a text cell a spreadsheet would read as a formula.

    Decides by the value's origin, before `fmt_number` renders anything: a number the tool
    produced is never bait, even after the comma locale turns it into printed text like "-12,50".
    Called on the raw value, ahead of fmt_number, for exactly that reason — call it after and a
    rendered negative number reads as a string starting with "-" and gets prefixed, which made a
    European refund into text.
    """
    if is_formula_bait(value):
        return "'" + value
    return value


def fmt_number(value: Any, cfg: Config) -> Any:
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)) and cfg.decimal == ",":
        return str(value).replace(".", ",")
    return value


def atomic_write(path: Path, write_fn) -> None:
    """Write to a temp file beside the target and rename, so an interrupted run leaves no half sheet."""
    tmp = path.with_name(path.name + ".part")
    try:
        write_fn(tmp)
        tmp.replace(path)
    finally:
        if tmp.exists():
            tmp.unlink()


def write_csv(path: Path, columns: list[str], rows: list[dict], trailer: dict, cfg: Config) -> None:
    def _w(target: Path) -> None:
        # UTF-8 BOM + CRLF is what makes Excel, Numbers and LibreOffice all open it clean.
        with target.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=cfg.separator, lineterminator="\r\n")
            w.writerow(headings(columns, header_style(cfg, "csv")))
            for r in rows:
                w.writerow([fmt_number(guard(r.get(c, "")), cfg) for c in columns])
            w.writerow([])
            w.writerow([f"# {k}: {v}" for k, v in trailer.items()])

    atomic_write(path, _w)


def write_json(path: Path, columns: list[str], rows: list[dict], trailer: dict, cfg: Config) -> None:
    """newline="" so Windows' text-mode translation never turns our own \\n into \\r\\n:
    bytes are the same on every platform, LF here, CRLF only in the csv writer."""
    keys = headings(columns, header_style(cfg, "json"))
    def _w(target: Path) -> None:
        payload = {"rows": [{k: r.get(c, "") for k, c in zip(keys, columns)} for r in rows],
                   "trailer": trailer}
        with target.open("w", newline="", encoding="utf-8") as fh:
            fh.write(json.dumps(payload, indent=2, ensure_ascii=False))

    atomic_write(path, _w)


def write_jsonl(path: Path, columns: list[str], rows: list[dict], cfg: Config) -> None:
    keys = headings(columns, header_style(cfg, "jsonl"))
    def _w(target: Path) -> None:
        with target.open("w", newline="", encoding="utf-8") as fh:
            for r in rows:
                obj = {k: r.get(c, "") for k, c in zip(keys, columns)}
                fh.write(json.dumps(obj, ensure_ascii=False) + "\n")

    atomic_write(path, _w)


def _force_text(cells) -> None:
    """XLSX holds a real cell type, so bait is stored as text instead of being prefixed."""
    for cell in cells:
        if is_formula_bait(cell.value):
            # openpyxl types a leading '=' as a formula; a plain assignment retypes the cell
            # (set_explicit_value is gone in 3.1). The file then carries t="inlineStr".
            cell.data_type = "s"


def write_xlsx(path: Path, columns: list[str], rows: list[dict], trailer: dict, cfg: Config) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        raise ExportError("--format xlsx needs openpyxl: pip install openpyxl") from None

    def _w(target: Path) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Documents"
        heads = headings(columns, header_style(cfg, "xlsx"))
        ws.append(heads)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for r in rows:
            ws.append([r.get(c, "") for c in columns])
            _force_text(ws[ws.max_row])
        for i, (col, head) in enumerate(zip(columns, heads), start=1):
            width = max(len(str(head)), *(len(str(r.get(col, ""))) for r in rows)) if rows else len(str(head))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width + 2, 8), 60)
        ws2 = wb.create_sheet("Export info")
        for k, v in trailer.items():
            ws2.append([k, str(v)])
            _force_text(ws2[ws2.max_row])
        wb.save(target)

    atomic_write(path, _w)


def write_inventory(path: Path, rows: list[dict], trailer: dict) -> None:
    """Print-ready binder sheet ordered by ASN. Every value escaped - titles are untrusted."""
    def sort_key(r):
        asn = r.get("archive_serial_number")
        return (asn == "" or asn is None, asn if isinstance(asn, int) else 0)

    ordered = sorted(rows, key=sort_key)
    with_asn = [r for r in ordered if r.get("archive_serial_number") not in ("", None)]
    without = [r for r in ordered if r.get("archive_serial_number") in ("", None)]

    def table(items):
        out = []
        for r in items:
            out.append(
                "<tr><td class='asn'>{}</td><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
                    html.escape(str(r.get("archive_serial_number", "") or "")),
                    html.escape(str(r.get("title", ""))),
                    html.escape(str(r.get("correspondent", ""))),
                    html.escape(str(r.get("document_type", ""))),
                    html.escape(str(r.get("created", ""))),
                )
            )
        return "\n".join(out)

    def _w(target: Path) -> None:
        body = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Document inventory</title>
<style>
 body {{ font: 11pt/1.4 -apple-system, Segoe UI, Helvetica, Arial, sans-serif; color:#000; background:#fff; margin:24px; }}
 h1 {{ font-size:16pt; margin:0 0 4px; }} p.meta {{ color:#444; margin:0 0 16px; font-size:9pt; }}
 table {{ border-collapse:collapse; width:100%; }}
 th, td {{ border-bottom:1px solid #bbb; padding:4px 6px; text-align:left; vertical-align:top; }}
 th {{ border-bottom:2px solid #000; font-size:9pt; text-transform:uppercase; letter-spacing:.04em; }}
 td.asn {{ font-variant-numeric:tabular-nums; white-space:nowrap; }}
 h2 {{ font-size:12pt; margin:24px 0 6px; }}
 tr {{ page-break-inside:avoid; }} thead {{ display:table-header-group; }}
 @page {{ margin:15mm; }}
</style></head><body>
<h1>Document inventory</h1>
<p class="meta">{html.escape(str(trailer.get('rows written', '')))} documents &middot; {html.escape(str(trailer.get('generated', '')))}</p>
<table><thead><tr><th>ASN</th><th>Title</th><th>Correspondent</th><th>Type</th><th>Created</th></tr></thead>
<tbody>
{table(with_asn)}
</tbody></table>
{'<h2>No archive serial number</h2><table><thead><tr><th>ASN</th><th>Title</th><th>Correspondent</th><th>Type</th><th>Created</th></tr></thead><tbody>' + table(without) + '</tbody></table>' if without else ''}
</body></html>"""
        with target.open("w", newline="", encoding="utf-8") as fh:
            fh.write(body)

    atomic_write(path, _w)


# --------------------------------------------------------------------------------------
# pack
# --------------------------------------------------------------------------------------
SAFE_CHARS = re.compile(r"[^A-Za-z0-9._ -]")


def safe_filename(doc_id: int, title: str) -> str:
    """Build from the id plus a slugged title. Titles are untrusted - a title can be '../../x'."""
    t = unicodedata.normalize("NFKD", title or "")
    t = SAFE_CHARS.sub("_", t)
    t = re.sub(r"\.{2,}", "_", t)  # no '..' survives, even without a separator to use it
    t = re.sub(r"[\s_]+", " ", t).strip().strip(". ")
    t = t[:100].strip() or "document"
    return f"{doc_id}-{t}.pdf"


def pack_documents(client: Client, rows: list[dict], out_dir: Path, cfg: Config, sheet_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    resolved_root = out_dir.resolve()
    params = {"original": "true"} if cfg.pack_version == "original" else None
    for r in rows:
        did = r["id"]
        name = safe_filename(did, str(r.get("title", "")))
        target = (out_dir / name).resolve()
        # A title that tries to climb out must not.
        if not str(target).startswith(str(resolved_root) + os.sep):
            raise ExportError(f"refusing to write document {did} outside the pack directory")
        data = client.get_bytes(f"/api/documents/{did}/download/", params)
        if not data:
            raise ExportError(f"document {did} downloaded empty; not shipping an unverifiable packet")
        target.write_bytes(data)
        if target.stat().st_size == 0:
            raise ExportError(f"document {did} wrote zero bytes")
        # Forward slashes always. The sheet travels with the folder, and a packet built on
        # Windows has to resolve when the accountant opens it on a Mac.
        r["file"] = os.path.relpath(target, sheet_dir).replace(os.sep, "/")
        r["file:version"] = cfg.pack_version


# --------------------------------------------------------------------------------------
# sums
# --------------------------------------------------------------------------------------
def totals_row(rows: list[dict], cfg: Config, lk: Lookups, columns: list[str]) -> dict[str, Any]:
    total: dict[str, Any] = {c: "" for c in columns}
    total[columns[0]] = "TOTAL"
    for name in cfg.sums:
        matches = [f for f in lk.custom_fields if f["name"].lower() == name.lower()]
        if not matches:
            raise UsageError(f"--sum: no custom field named {name!r}")
        f = matches[0]
        if f["data_type"] not in SUMMABLE:
            raise UsageError(f"--sum: field {f['name']!r} is {f['data_type']}, not summable")
        col = lk.cf_column_name(f)
        if col not in columns:
            raise UsageError(f"--sum: field {f['name']!r} is not among the exported columns")
        if f["data_type"] == "monetary":
            currencies = {r.get(f"{col}:currency", "") for r in rows if r.get(col) not in ("", None)}
            currencies.discard("")
            if len(currencies) > 1:
                raise ExportError(
                    f"--sum: field {f['name']!r} mixes currencies {sorted(currencies)}. "
                    "Refusing a meaningless total."
                )
            if currencies:
                total[f"{col}:currency"] = currencies.pop()
        s = 0.0
        for r in rows:
            v = r.get(col)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                s += float(v)
        # An integer field totals as an integer; 5015.0 in a sheet is just wrong.
        total[col] = int(round(s)) if f["data_type"] == "integer" else round(s, 2)
    return total


# --------------------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="paperless-export",
        description="Export a paperless-ngx document list to a sheet you can check.",
        epilog="PAPERLESS_URL and PAPERLESS_TOKEN come from the environment. Never pass a token as a flag.",
    )
    p.add_argument("--view", help="saved view id or name; its filters, sort and columns win")
    p.add_argument("--query", help="full text search")
    p.add_argument("--title", help="title contains")
    p.add_argument("--correspondent", action="append", default=[])
    p.add_argument("--document-type", action="append", default=[], dest="document_type")
    p.add_argument("--tag", action="append", default=[], dest="tags")
    p.add_argument("--storage-path", action="append", default=[], dest="storage_path")
    p.add_argument("--created-from", dest="created_from", help="YYYY-MM-DD")
    p.add_argument("--created-to", dest="created_to", help="YYYY-MM-DD")
    p.add_argument("--custom-field-query", dest="custom_field_query", help="paperless custom_field_query JSON")
    p.add_argument("--sum", action="append", default=[], dest="sums", metavar="FIELD")
    p.add_argument("--columns", help="comma separated column list")
    p.add_argument("--content", action="store_true", help="include the full OCR text column")
    p.add_argument("--locale", choices=["dot", "comma"], default="dot",
                   help="dot: ',' separator '.' decimal (default). comma: ';' separator ',' decimal")
    p.add_argument("--date", default="iso", dest="date",
                   help="iso (default), us (MM/DD/YYYY), eu (DD.MM.YYYY), or a strftime pattern")
    p.add_argument("--headers", choices=["readable", "raw"], default=None,
                   help="column headings: readable (default for csv and xlsx) or raw (default for json and jsonl)")
    p.add_argument("--list-sep", default="; ", dest="list_sep",
                   help="joins multi-value cells such as tags. Default '; '")
    p.add_argument("-o", "--output", default="export", help="output path without extension")
    p.add_argument("--format", action="append", choices=["csv", "json", "jsonl", "xlsx"],
                   default=[], dest="formats")
    p.add_argument("--pack", action="store_true", help="also download the matching PDFs beside the sheet")
    p.add_argument("--pack-version", choices=["latest", "original"], default="latest", dest="pack_version")
    p.add_argument("--inventory", action="store_true", help="also write a print-ready HTML binder sheet")
    p.add_argument("--insecure", action="store_true", help="allow plain http to a public host")
    # Front 3. --serve does not export; it serves the page on 127.0.0.1 for one session.
    # There is no --host and no --port, and no flag will add them: the bind address is fixed.
    p.add_argument("--serve", action="store_true",
                   help="serve the page on 127.0.0.1 for one session instead of exporting now")
    p.add_argument("--idle", type=int, default=900, metavar="SECONDS",
                   help="--serve only: shut down after this long with no request (default 900, min 60)")
    p.add_argument("--no-open", action="store_true", dest="no_open",
                   help="--serve only: print the URL instead of opening a browser")
    # Answered during parsing, so it works with no environment and no instance: someone who
    # cannot run an export at all can still tell us which build they have.
    p.add_argument("--version", action="version", version=f"paperless-export {__version__}")
    return p


def config_from_args(argv: list[str]) -> Config:
    args = build_parser().parse_args(argv)
    url = os.environ.get("PAPERLESS_URL", "").strip()
    token = os.environ.get("PAPERLESS_TOKEN", "").strip()
    if not url or not token:
        raise UsageError("PAPERLESS_URL and PAPERLESS_TOKEN must be set in the environment")
    _SECRETS.append(token)
    if args.list_sep == (";" if args.locale == "comma" else ","):
        raise UsageError(
            f"--list-sep {args.list_sep!r} is the CSV separator this locale uses. "
            "Quoting would make it legal CSV and half the readers would still split the cell."
        )
    if args.view and any([args.query, args.title, args.correspondent, args.document_type,
                          args.tags, args.storage_path, args.created_from, args.created_to,
                          args.custom_field_query]):
        raise UsageError("--view carries its own filters; do not combine it with filter flags")
    # Fail before a single file is written, not halfway through the format list.
    if "xlsx" in (args.formats or []):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            raise UsageError("--format xlsx needs openpyxl: pip install openpyxl") from None

    return Config(
        base_url=check_url(url, args.insecure),
        token=token,
        view=args.view,
        filters={
            "query": args.query, "title": args.title,
            "correspondent": args.correspondent, "document_type": args.document_type,
            "tags": args.tags, "storage_path": args.storage_path,
            "created_from": args.created_from, "created_to": args.created_to,
            "custom_field_query": args.custom_field_query,
        },
        sums=args.sums,
        columns=[c.strip() for c in args.columns.split(",")] if args.columns else None,
        content=args.content,
        separator=";" if args.locale == "comma" else ",",
        decimal="," if args.locale == "comma" else ".",
        date_spec=args.date,
        date_fmt=resolve_date_format(args.date),
        headers=args.headers,
        list_sep=args.list_sep,
        output=Path(args.output),
        formats=args.formats or ["csv"],
        pack=args.pack,
        pack_version=args.pack_version,
        inventory=args.inventory,
        insecure=args.insecure,
    )


# --------------------------------------------------------------------------------------
# run
# --------------------------------------------------------------------------------------
def export(cfg: Config, emit=None) -> dict:
    """Run the export. Returns the trailer and the files written; `emit` gets the progress lines.

    The CLI and the loopback front both come through here, so neither can drift from the other.
    """
    emit = emit or (lambda _msg: None)
    written: list[Path] = []
    started = time.time()
    notes: list[str] = []
    client = Client(cfg.base_url, cfg.token)
    lk = Lookups(client)

    view_columns = None
    sort = None
    if cfg.view:
        views = client.get_all("/api/saved_views/")
        match = [v for v in views if str(v["id"]) == str(cfg.view)]
        if not match:
            match = [v for v in views if v["name"] == cfg.view]
        if not match:
            match = [v for v in views if v["name"].lower() == str(cfg.view).lower()]
        if not match:
            names = ", ".join(sorted(v["name"] for v in views)) or "(none)"
            raise UsageError(f"no saved view {cfg.view!r}. Available: {names}")
        if len(match) > 1:
            raise UsageError(f"saved view {cfg.view!r} is ambiguous (ids {[v['id'] for v in match]})")
        params, sort, view_columns = params_from_view(match[0], notes)
    else:
        params = params_from_flags(cfg, lk)

    if sort:
        params["ordering"] = sort

    docs, api_count = fetch_documents(client, params, cfg)
    columns = resolve_columns(cfg, lk, view_columns, notes)
    rows = [build_row(d, lk, cfg, notes) for d in docs]

    if any(header_style(cfg, f) == "readable" for f in cfg.formats) or cfg.inventory:
        clash = readable_clashes(columns)
        if clash:
            notes.append(f"readable headings {clash} would collide; those columns keep their raw names")

    # Roots only. The list endpoint guarantees this server-side; assert it rather than assume.
    strays = [r["id"] for r, d in zip(rows, docs) if d.get("root_document") is not None]
    if strays:
        raise ExportError(f"documents {strays} are versions, not roots; refusing to export them as rows")

    # Reconcile before packing: never download a thousand PDFs for an export we already know
    # we cannot vouch for.
    if len(rows) != api_count:
        raise ReconcileError(
            f"rows written {len(rows)} != api count {api_count}. The export is not trustworthy."
        )

    sheet_dir = cfg.output.parent.resolve()
    sheet_dir.mkdir(parents=True, exist_ok=True)
    packed: list[Path] = []
    if cfg.pack:
        pack_dir = cfg.output.parent / f"{cfg.output.name}-{dt.date.today().isoformat()}"
        pack_documents(client, rows, pack_dir, cfg, sheet_dir)
        packed = sorted(q for q in pack_dir.rglob("*") if q.is_file())

    out_rows = list(rows)
    if cfg.sums:
        out_rows.append(totals_row(rows, cfg, lk, columns))

    trailer = {
        "rows written": len(rows),
        "api count": api_count,
        "filter": urllib.parse.urlencode(params) or "(none)",
        "formats": ", ".join(cfg.formats),
        "headers": "",  # replaced per file below; here to hold its place in the order
        "date format": cfg.date_spec,
        "list separator": repr(cfg.list_sep),
        "generated": dt.datetime.now().isoformat(timespec="seconds"),
        "elapsed": f"{time.time() - started:.1f}s",
    }
    if notes:
        trailer["notes"] = " | ".join(dict.fromkeys(notes))

    for fmt in cfg.formats:
        target = cfg.output.with_suffix(f".{fmt}")
        # Header style differs per format, so each file states its own.
        per_file = dict(trailer, headers=header_style(cfg, fmt))
        if fmt == "csv":
            write_csv(target, columns, out_rows, per_file, cfg)
        elif fmt == "json":
            write_json(target, columns, out_rows, per_file, cfg)
        elif fmt == "jsonl":
            write_jsonl(target, columns, out_rows, cfg)
        elif fmt == "xlsx":
            write_xlsx(target, columns, out_rows, per_file, cfg)
        written.append(target)
        emit(f"wrote {target}")

    if cfg.inventory:
        inv = cfg.output.with_suffix(".inventory.html")
        write_inventory(inv, rows, trailer)
        written.append(inv)
        emit(f"wrote {inv}")

    return {
        "trailer": dict(trailer, headers=header_style(cfg, cfg.formats[0])),
        "files": written,
        "packed": packed,
        "notes": notes,
    }


def run(cfg: Config) -> int:
    result = export(cfg, emit=print)
    print()
    for k, v in result["trailer"].items():
        print(f"{k:<15}: {v}")
    return EXIT_OK


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    try:
        if build_parser().parse_args(argv).serve:
            # Imported here, not at the top: the CLI must not carry the server on every run.
            import paperless_serve
            return paperless_serve.serve(argv)
        return run(config_from_args(argv))
    except UsageError as e:
        print(redact(f"usage error: {e}"), file=sys.stderr)
        return EXIT_USAGE
    except ReconcileError as e:
        print(redact(f"RECONCILIATION FAILED: {e}"), file=sys.stderr)
        return EXIT_RECONCILE
    except ExportError as e:
        print(redact(f"export failed: {e}"), file=sys.stderr)
        return EXIT_FAIL
    except KeyboardInterrupt:
        print("interrupted", file=sys.stderr)
        return EXIT_FAIL


if __name__ == "__main__":
    sys.exit(main())
