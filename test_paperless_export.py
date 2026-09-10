#!/usr/bin/env python3
"""Tests for paperless_export. Stdlib unittest; no live server needed.

Run: python3 -m unittest discover -s execution -v
The live run against a real instance is the other half of verification; this suite is not the
whole of it.
"""
from __future__ import annotations

import ast
import contextlib
import csv
import io
import json
import re
import shutil
import sys
import tempfile
import tokenize
import unittest
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import gen_rules              # noqa: E402
import paperless_export as px  # noqa: E402

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False


# --------------------------------------------------------------------------------------
# recorded API fixture: shaped exactly like v3.1.3 / API v10 returns
# --------------------------------------------------------------------------------------
# One fixture for both engines. The JavaScript conformance run reads the same file, so the two
# engines can never be tested against different paperlesses.
FIXTURE = json.loads((Path(__file__).resolve().parent / "fixture_v10.json").read_text(encoding="utf-8"))
CUSTOM_FIELDS = FIXTURE["custom_fields"]
DOCS = FIXTURE["documents"]
SAVED_VIEWS = FIXTURE["saved_views"]


class FakeClient:
    """Stands in for Client. Returns the fixture; records nothing is ever written."""

    def __init__(self, docs=None, count=None):
        self.base = "https://paperless.test"
        self._docs = DOCS if docs is None else docs
        self._count = len(self._docs) if count is None else count

    def get_all(self, path):
        return {
            "/api/correspondents/": [{"id": 1, "name": "Acme"}],
            "/api/document_types/": [{"id": 1, "name": "Invoice"}],
            "/api/tags/": [{"id": 1, "name": "tax"}, {"id": 2, "name": "2025"}],
            "/api/storage_paths/": [{"id": 1, "name": "Taxes"}],
            "/api/custom_fields/": CUSTOM_FIELDS,
            "/api/saved_views/": SAVED_VIEWS,
        }[path]

    def get(self, path, params=None):
        return {"count": self._count, "next": None, "results": self._docs}, {}

    def get_bytes(self, path, params=None):
        return b"%PDF-1.4 fake"


def make_cfg(**kw):
    base = dict(base_url="https://paperless.test", token="t")
    base.update(kw)
    return px.Config(**base)


def lookups():
    return px.Lookups(FakeClient())


@contextlib.contextmanager
def run_with(client):
    """Swap in the fixture client and swallow the run's own output."""
    original = px.Client
    px.Client = lambda *a, **k: client
    try:
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            yield
    finally:
        px.Client = original


# --------------------------------------------------------------------------------------
class TestRuleMap(unittest.TestCase):
    def test_every_rule_type_0_to_50_is_mapped(self):
        missing = [i for i in range(51) if i not in px.RULE_MAP]
        self.assertEqual(missing, [], f"unmapped saved-view rule types: {missing}")

    def test_known_mappings(self):
        self.assertEqual(px.RULE_MAP[6][0], "tags__id__all")
        self.assertEqual(px.RULE_MAP[22][0], "tags__id__in")
        self.assertEqual(px.RULE_MAP[42][0], "custom_field_query")
        self.assertEqual(px.RULE_MAP[3][1], "correspondent__isnull")


class TestViewParams(unittest.TestCase):
    def test_multi_values_join_with_commas(self):
        view = {"filter_rules": [{"rule_type": 6, "value": "1"}, {"rule_type": 6, "value": "2"}]}
        params, _, _ = px.params_from_view(view, [])
        self.assertEqual(params["tags__id__all"], "1,2")

    def test_null_value_uses_isnull_param(self):
        view = {"filter_rules": [{"rule_type": 3, "value": None}]}
        params, _, _ = px.params_from_view(view, [])
        self.assertEqual(params["correspondent__isnull"], "true")

    def test_sort_reverse_prefixes_minus(self):
        view = {"filter_rules": [], "sort_field": "created", "sort_reverse": True}
        _, sort, _ = px.params_from_view(view, [])
        self.assertEqual(sort, "-created")

    def test_unknown_rule_type_is_a_hard_error(self):
        """A silently dropped filter yields too many rows that still reconcile. Refuse instead."""
        view = {"filter_rules": [{"rule_type": 999, "value": "x"}]}
        with self.assertRaises(px.ExportError) as ctx:
            px.params_from_view(view, [])
        self.assertIn("999", str(ctx.exception))


class TestUrlPolicy(unittest.TestCase):
    """Resolution is always stubbed here: these tests must not depend on a working DNS."""

    def setUp(self):
        self._real_resolve = px._resolve_host
        px._resolve_host = lambda host: self.dns.get(host, [])
        self.dns: dict[str, list[str]] = {}

    def tearDown(self):
        px._resolve_host = self._real_resolve

    def test_bare_lan_name_resolving_private_is_private(self):
        self.dns["nas"] = ["192.168.1.20"]
        self.assertTrue(px.is_private_host("nas"))
        self.assertEqual(px.check_url("http://nas:8000", False), "http://nas:8000")

    def test_home_arpa_name_resolving_private_is_private(self):
        self.dns["paperless.home.arpa"] = ["10.0.0.9", "fd00::5"]
        self.assertTrue(px.is_private_host("paperless.home.arpa"))
        self.assertEqual(
            px.check_url("http://paperless.home.arpa", False), "http://paperless.home.arpa"
        )

    def test_public_name_is_not_private(self):
        self.dns["p.example.com"] = ["93.184.216.34"]
        self.assertFalse(px.is_private_host("p.example.com"))

    def test_mixed_resolution_is_not_private(self):
        """One public address is enough: the token would leave the LAN on that route."""
        self.dns["split.example"] = ["192.168.1.20", "93.184.216.34"]
        self.assertFalse(px.is_private_host("split.example"))
        with self.assertRaises(px.UsageError):
            px.check_url("http://split.example", False)

    def test_unresolvable_name_is_not_private(self):
        self.assertFalse(px.is_private_host("nowhere.invalid"))
        with self.assertRaises(px.UsageError):
            px.check_url("http://nowhere.invalid", False)

    def test_shortcuts_need_no_resolution(self):
        for host in ("localhost", "nas.local", "box.localhost", "NAS.LOCAL", "nas.local."):
            self.assertTrue(px.is_private_host(host), host)

    def test_public_literal_ip_is_never_resolved(self):
        """A literal address is judged as itself; a reverse lookup must not launder it."""
        self.dns["93.184.216.34"] = ["192.168.1.20"]
        self.assertFalse(px.is_private_host("93.184.216.34"))

    def test_ipv6_private_forms(self):
        self.assertTrue(px.is_private_host("::1"))
        self.assertTrue(px.is_private_host("fd00::5"))
        self.assertTrue(px.is_private_host("fe80::1%en0"))
        self.assertFalse(px.is_private_host("2606:4700::1111"))

    def test_ipv4_mapped_ipv6_is_judged_on_the_mapped_address(self):
        self.assertTrue(px.is_private_host("::ffff:192.168.1.20"))
        self.assertFalse(px.is_private_host("::ffff:93.184.216.34"))

    def test_https_allowed(self):
        self.assertEqual(px.check_url("https://p.example.com", False), "https://p.example.com")

    def test_http_to_public_host_refused(self):
        self.dns["p.example.com"] = ["93.184.216.34"]
        with self.assertRaises(px.UsageError):
            px.check_url("http://p.example.com", False)

    def test_http_to_public_host_allowed_with_insecure(self):
        self.dns["p.example.com"] = ["93.184.216.34"]
        self.assertEqual(px.check_url("http://p.example.com", True), "http://p.example.com")

    def test_http_to_private_host_allowed(self):
        for host in ("http://127.0.0.1:8000", "http://192.168.1.10", "http://10.0.0.5",
                     "http://nas.local", "http://localhost:8000"):
            self.assertTrue(px.check_url(host, False).startswith("http://"))

    def test_url_with_query_refused(self):
        with self.assertRaises(px.UsageError):
            px.check_url("https://p.example.com/?token=secret", False)

    def test_bad_scheme_refused(self):
        with self.assertRaises(px.UsageError):
            px.check_url("ftp://p.example.com", False)


class TestRedaction(unittest.TestCase):
    def setUp(self):
        px._SECRETS.append("sekrit-token-value")

    def tearDown(self):
        px._SECRETS.remove("sekrit-token-value")

    def test_a_token_echoed_by_the_server_never_reaches_the_terminal(self):
        body = "export failed: /api/documents/ returned HTTP 400: Token sekrit-token-value bad"
        out = px.redact(body)
        self.assertNotIn("sekrit-token-value", out)
        self.assertIn("***REDACTED***", out)

    def test_ordinary_text_untouched(self):
        self.assertEqual(px.redact("nothing to see"), "nothing to see")

    def test_a_trivially_short_secret_does_not_shred_the_message(self):
        px._SECRETS.append("t")
        try:
            self.assertEqual(px.redact("nothing to see"), "nothing to see")
        finally:
            px._SECRETS.remove("t")


class TestRedirectPolicy(unittest.TestCase):
    """urllib carries Authorization across a redirect. So the redirect must not leave home."""

    def redirect(self, frm, to):
        h = px.SameOriginRedirect()
        req = urllib.request.Request(frm)
        return h.redirect_request(req, None, 302, "Found", {}, to)

    def test_same_origin_redirect_allowed(self):
        r = self.redirect("https://p.example.com/api/documents/",
                          "https://p.example.com/api/documents/?page=2")
        self.assertIsNotNone(r)

    def test_https_to_http_refused(self):
        with self.assertRaises(px.ExportError) as ctx:
            self.redirect("https://p.example.com/api/", "http://p.example.com/api/")
        self.assertIn("https->http", str(ctx.exception))

    def test_cross_host_redirect_refused(self):
        with self.assertRaises(px.ExportError):
            self.redirect("https://p.example.com/api/", "https://evil.example/api/")

    def test_cross_port_redirect_refused(self):
        with self.assertRaises(px.ExportError):
            self.redirect("http://192.168.1.5:8000/api/", "http://192.168.1.5:9000/api/")

    def test_default_port_is_not_a_change_of_origin(self):
        r = self.redirect("https://p.example.com/api/", "https://p.example.com:443/api/x")
        self.assertIsNotNone(r)


class TestPaginationOrigin(unittest.TestCase):
    """paperless builds `next` from its own configured hostname, not the one the user typed."""

    def client(self, base="http://192.168.1.5:8000"):
        return px.Client(base, "tok")

    def test_relative_next_becomes_a_path(self):
        c = self.client()
        self.assertEqual(c.next_path("/api/documents/?page=2"), "/api/documents/?page=2")

    def test_same_origin_absolute_next_is_reduced_to_a_path(self):
        c = self.client()
        self.assertEqual(
            c.next_path("http://192.168.1.5:8000/api/documents/?page=2&page_size=1000"),
            "/api/documents/?page=2&page_size=1000",
        )

    def test_next_on_another_origin_is_refused(self):
        c = self.client()
        with self.assertRaises(px.ExportError):
            c.next_path("http://paperless.lan/api/documents/?page=2")

    def test_old_behaviour_would_have_mangled_the_url(self):
        """The regression this replaces: a differently-named `next` was concatenated onto base."""
        c = self.client()
        nxt = "http://paperless.lan/api/documents/?page=2"
        self.assertFalse(nxt.startswith(c.base))  # the condition the old code used
        with self.assertRaises(px.ExportError):
            c.next_path(nxt)


class TestInjectionGuard(unittest.TestCase):
    def test_formula_prefixes_are_escaped(self):
        for bad in ("=1+1", "+1", "-1", "@SUM(1)", "\tx", "\rx"):
            self.assertTrue(px.guard(bad).startswith("'"), f"{bad!r} not guarded")

    def test_ordinary_text_untouched(self):
        self.assertEqual(px.guard("Invoice 2025"), "Invoice 2025")

    def test_numbers_pass_through_unchanged(self):
        self.assertEqual(px.guard(-5), -5)
        self.assertEqual(px.guard(12.5), 12.5)


class TestSafeFilename(unittest.TestCase):
    def test_path_traversal_is_flattened(self):
        name = px.safe_filename(9, "../../../../tmp/pwned")
        self.assertNotIn("/", name)
        self.assertNotIn("..", name)
        self.assertTrue(name.startswith("9-"))

    def test_no_separators_survive(self):
        for hostile in ("a/b", "a\\b", "../x", "..", "."):
            name = px.safe_filename(1, hostile)
            self.assertNotIn("/", name)
            self.assertNotIn("\\", name)

    def test_empty_title_still_produces_a_name(self):
        self.assertEqual(px.safe_filename(3, ""), "3-document.pdf")

    def test_length_is_capped(self):
        self.assertLessEqual(len(px.safe_filename(1, "x" * 500)), 110)


class TestPackPaths(unittest.TestCase):
    def test_file_column_uses_forward_slashes(self):
        """A packet built on Windows must resolve when it is opened on a Mac."""
        import os
        with tempfile.TemporaryDirectory() as td:
            sheet_dir = Path(td).resolve()  # run() resolves it too; /var is a symlink on macOS
            pack_dir = sheet_dir / "out-2026-01-01"
            rows = [{"id": 1, "title": "Invoice"}]
            cfg = make_cfg(output=sheet_dir / "out", pack=True)
            px.pack_documents(FakeClient(), rows, pack_dir, cfg, sheet_dir)
            self.assertNotIn("\\", rows[0]["file"])
            self.assertEqual(rows[0]["file"], "out-2026-01-01/1-Invoice.pdf")
            self.assertEqual(rows[0]["file:version"], "latest")
            del os


class TestMonetary(unittest.TestCase):
    def test_currency_and_value_split(self):
        self.assertEqual(px.parse_monetary("USD123.45"), ("USD", 123.45))
        self.assertEqual(px.parse_monetary("EUR12.50"), ("EUR", 12.5))

    def test_bare_number(self):
        self.assertEqual(px.parse_monetary("10.00"), ("", 10.0))

    def test_empty(self):
        self.assertEqual(px.parse_monetary(None), ("", None))
        self.assertEqual(px.parse_monetary(""), ("", None))


class TestFlatten(unittest.TestCase):
    def test_select_resolves_to_label_not_option_id(self):
        lk = lookups()
        row = px.flatten_custom_fields(DOCS[0], lk, make_cfg(), [])
        self.assertEqual(row["custom_field:Category"], "Utilities")

    def test_unknown_select_option_written_verbatim_and_flagged(self):
        lk = lookups()
        notes = []
        doc = dict(DOCS[0], custom_fields=[{"field": 7, "value": "nope"}])
        row = px.flatten_custom_fields(doc, lk, make_cfg(), notes)
        self.assertEqual(row["custom_field:Category"], "nope")
        self.assertTrue(any("no label" in n for n in notes))

    def test_monetary_splits_into_value_and_currency(self):
        lk = lookups()
        row = px.flatten_custom_fields(DOCS[0], lk, make_cfg(), [])
        self.assertEqual(row["custom_field:Amount"], 10.5)
        self.assertEqual(row["custom_field:Amount:currency"], "USD")

    def test_documentlink_joins_ids(self):
        lk = lookups()
        row = px.flatten_custom_fields(DOCS[0], lk, make_cfg(), [])
        self.assertEqual(row["custom_field:Related Doc"], "2")

    def test_absent_field_is_empty_not_zero(self):
        lk = lookups()
        row = px.flatten_custom_fields(DOCS[1], lk, make_cfg(), [])
        self.assertEqual(row["custom_field:Category"], "")
        self.assertEqual(row["custom_field:Reviewed"], "")


class TestColumnOrder(unittest.TestCase):
    def test_custom_fields_are_alphabetical_by_name(self):
        lk = lookups()
        cols = px.resolve_columns(make_cfg(), lk, None, [])
        cf = [c for c in cols if c.startswith("custom_field:") and not c.endswith(":currency")]
        self.assertEqual(cf, sorted(cf), "custom field columns must be alphabetical by name")

    def test_id_first_url_last(self):
        cols = px.resolve_columns(make_cfg(), lookups(), None, [])
        self.assertEqual(cols[0], "id")
        self.assertEqual(cols[-1], "url")


class TestTotals(unittest.TestCase):
    def test_integer_field_totals_as_an_integer(self):
        lk = lookups()
        cfg = make_cfg(sums=["Invoice Number"])
        cols = px.resolve_columns(cfg, lk, None, [])
        rows = [px.build_row(d, lk, cfg, []) for d in DOCS]
        total = px.totals_row(rows, cfg, lk, cols)
        self.assertEqual(total["custom_field:Invoice Number"], 300)
        self.assertNotIsInstance(total["custom_field:Invoice Number"], float)

    def test_monetary_total_and_currency(self):
        lk = lookups()
        cfg = make_cfg(sums=["Amount"])
        cols = px.resolve_columns(cfg, lk, None, [])
        rows = [px.build_row(d, lk, cfg, []) for d in DOCS]
        total = px.totals_row(rows, cfg, lk, cols)
        self.assertEqual(total["custom_field:Amount"], 15.75)
        self.assertEqual(total["custom_field:Amount:currency"], "USD")

    def test_mixed_currency_is_refused(self):
        lk = lookups()
        cfg = make_cfg(sums=["Amount"])
        cols = px.resolve_columns(cfg, lk, None, [])
        docs = [dict(DOCS[0]), dict(DOCS[1], custom_fields=[{"field": 4, "value": "EUR5.00"}])]
        rows = [px.build_row(d, lk, cfg, []) for d in docs]
        with self.assertRaises(px.ExportError) as ctx:
            px.totals_row(rows, cfg, lk, cols)
        self.assertIn("currenc", str(ctx.exception).lower())

    def test_non_summable_type_refused(self):
        lk = lookups()
        cfg = make_cfg(sums=["Category"])
        cols = px.resolve_columns(cfg, lk, None, [])
        with self.assertRaises(px.UsageError):
            px.totals_row([], cfg, lk, cols)


class TestReconciliation(unittest.TestCase):
    def test_mismatch_raises_and_exits_3(self):
        """The core promise: never present a sheet the tool cannot check."""
        with tempfile.TemporaryDirectory() as td:
            cfg = make_cfg(output=Path(td) / "out")
            with run_with(FakeClient(count=99)):  # server says 99, we got 2
                with self.assertRaises(px.ReconcileError):
                    px.run(cfg)

    def test_mismatch_exits_3_through_main(self):
        import os
        with tempfile.TemporaryDirectory() as td:
            old = {k: os.environ.get(k) for k in ("PAPERLESS_URL", "PAPERLESS_TOKEN")}
            os.environ["PAPERLESS_URL"] = "https://paperless.test"
            os.environ["PAPERLESS_TOKEN"] = "t"
            try:
                with run_with(FakeClient(count=99)):
                    rc = px.main(["-o", str(Path(td) / "out")])
            finally:
                for k, v in old.items():
                    if v is None:
                        os.environ.pop(k, None)
                    else:
                        os.environ[k] = v
            self.assertEqual(rc, px.EXIT_RECONCILE)

    def test_roots_only_enforced(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = make_cfg(output=Path(td) / "out")
            with run_with(FakeClient(docs=[dict(DOCS[0], root_document=7)])):
                with self.assertRaises(px.ExportError) as ctx:
                    px.run(cfg)
            self.assertIn("versions, not roots", str(ctx.exception))


class TestWriters(unittest.TestCase):
    def _run(self, **kw):
        td = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, td, ignore_errors=True)
        cfg = make_cfg(output=Path(td) / "out", **kw)
        with run_with(FakeClient()):
            rc = px.run(cfg)
        return rc, Path(td)

    def test_csv_has_bom_and_crlf(self):
        rc, td = self._run(formats=["csv"])
        self.assertEqual(rc, 0)
        raw = (td / "out.csv").read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "CSV needs a UTF-8 BOM for Excel")
        self.assertIn(b"\r\n", raw, "CSV needs CRLF")

    def test_csv_guards_the_formula_title(self):
        rc, td = self._run(formats=["csv"])
        with (td / "out.csv").open(encoding="utf-8-sig") as fh:
            rows = list(csv.reader(fh))
        titles = [r[rows[0].index("Title")] for r in rows[1:] if len(r) > 1]
        self.assertIn("'=1+1", titles)

    def test_json_carries_the_trailer(self):
        rc, td = self._run(formats=["json"])
        payload = json.loads((td / "out.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["trailer"]["rows written"], 2)
        self.assertEqual(payload["trailer"]["api count"], 2)

    def test_jsonl_one_object_per_line(self):
        rc, td = self._run(formats=["jsonl"])
        lines = [l for l in (td / "out.jsonl").read_text(encoding="utf-8").splitlines() if l.strip()]
        self.assertEqual(len(lines), 2)
        for l in lines:
            json.loads(l)

    def test_json_jsonl_and_inventory_are_lf_only_on_every_platform(self):
        """newline="" on every text writer means Windows cannot turn our \\n into \\r\\n. The gap
        was present since 0.1.0 for json, jsonl and the inventory; the kit had only ever compared
        csv, which is why it went unnoticed until a byte-for-byte compare covered the others.
        """
        rc, td = self._run(formats=["json", "jsonl"], inventory=True)
        for name in ("out.json", "out.jsonl", "out.inventory.html"):
            raw = (td / name).read_bytes()
            self.assertNotIn(b"\r\n", raw, f"{name} must be LF-only, not CRLF")
            self.assertIn(b"\n", raw, f"{name} has no line breaks to check")

    def test_locale_comma_uses_semicolon_separator(self):
        rc, td = self._run(formats=["csv"], separator=";", decimal=",")
        head = (td / "out.csv").read_text(encoding="utf-8-sig").splitlines()[0]
        self.assertIn(";", head)

    def test_a_rendered_negative_number_is_never_guarded_even_in_the_comma_locale(self):
        """The guard decides by origin, not by the rendered first character. Before this was fixed,
        fmt_number ran first, turned -12.5 into the string "-12,5", and guard then saw a leading
        "-" and prefixed it — a European refund coming out as text.
        """
        td = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, td, ignore_errors=True)
        cfg = make_cfg(output=Path(td) / "out", formats=["csv"], separator=";", decimal=",")
        docs = [d for d in FIXTURE["hostile_documents"] if d["id"] == 93]
        with run_with(FakeClient(docs=docs)):
            rc = px.run(cfg)
        self.assertEqual(rc, 0)
        with (Path(td) / "out.csv").open(encoding="utf-8-sig") as fh:
            rows = list(csv.reader(fh, delimiter=";"))
        header = rows[0]
        cell = rows[1][header.index("Amount")]
        self.assertEqual(cell, "-12,5")
        self.assertFalse(cell.startswith("'"), f"a rendered number was guarded: {cell!r}")

    def test_inventory_escapes_html(self):
        td = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, td, ignore_errors=True)
        rows = [{"archive_serial_number": 1, "title": "<script>alert(1)</script>",
                 "correspondent": "a", "document_type": "b", "created": "2025-01-01"}]
        px.write_inventory(td / "inv.html", rows, {"rows written": 1, "generated": "now"})
        body = (td / "inv.html").read_text(encoding="utf-8")
        self.assertNotIn("<script>alert(1)</script>", body)
        self.assertIn("&lt;script&gt;", body)

    def test_no_part_file_left_behind(self):
        rc, td = self._run(formats=["csv"])
        self.assertEqual(list(td.glob("*.part")), [])

    @unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
    def test_xlsx_keeps_numbers_numeric(self):
        rc, td = self._run(formats=["xlsx"])
        from openpyxl import load_workbook
        wb = load_workbook(td / "out.xlsx")
        ws = wb["Documents"]
        hdr = [c.value for c in ws[1]]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        amounts = [r[hdr.index("Amount")] for r in rows]
        self.assertIn(10.5, amounts, "monetary must stay a number so totals work")
        self.assertIn("Export info", wb.sheetnames)

    @unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
    def test_xlsx_writes_bait_as_a_text_cell_with_no_prefix(self):
        """XLSX guards by cell type, not by an apostrophe: the stored value stays verbatim."""
        rc, td = self._run(formats=["xlsx"])
        from openpyxl import load_workbook
        wb = load_workbook(td / "out.xlsx")
        ws = wb["Documents"]
        hdr = [c.value for c in ws[1]]
        col = hdr.index("Title") + 1
        titles = {ws.cell(row=r, column=col).value for r in range(2, ws.max_row + 1)}
        self.assertIn("=1+1", titles, "the title must survive verbatim")
        self.assertNotIn("'=1+1", titles, "no apostrophe may be stored in an xlsx cell")
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=col)
            if cell.value == "=1+1":
                self.assertEqual(cell.data_type, "s", "bait must be an explicit string cell")

    @unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
    def test_xlsx_all_six_prefixes_are_string_cells(self):
        from openpyxl import load_workbook
        bait = ["=1+1", "+1", "-1", "@SUM(1)", "\tx", "\rx"]
        rows = [{"title": b} for b in bait]
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "b.xlsx"
            px.write_xlsx(out, ["title"], rows, {"rows written": len(bait), "note": "=EVIL()"},
                          make_cfg())
            wb = load_workbook(out)
            ws = wb["Documents"]
            for i, b in enumerate(bait, start=2):
                cell = ws.cell(row=i, column=1)
                # XML normalises a literal CR to LF on read; the cell is still text, which is
                # what the guard is for. Every other character survives byte for byte.
                expected = b.replace("\r", "\n")
                self.assertEqual(cell.value, expected, f"{b!r} altered")
                self.assertEqual(cell.data_type, "s", f"{b!r} is not a string cell")
                self.assertNotEqual(cell.data_type, "f", f"{b!r} stored as a formula")
            info = wb["Export info"]
            trailer_cell = [c for row in info.iter_rows() for c in row if c.value == "=EVIL()"]
            self.assertTrue(trailer_cell, "trailer value missing")
            self.assertEqual(trailer_cell[0].data_type, "s", "the Export info sheet is guarded too")

    @unittest.skipIf(HAVE_OPENPYXL, "openpyxl is installed")
    def test_xlsx_without_openpyxl_fails_before_writing_anything(self):
        import os
        with tempfile.TemporaryDirectory() as td:
            old = {k: os.environ.get(k) for k in ("PAPERLESS_URL", "PAPERLESS_TOKEN")}
            os.environ["PAPERLESS_URL"] = "https://paperless.test"
            os.environ["PAPERLESS_TOKEN"] = "t"
            try:
                with run_with(FakeClient()):
                    rc = px.main(["--format", "csv", "--format", "xlsx", "-o", str(Path(td) / "out")])
            finally:
                for k, v in old.items():
                    os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)
            self.assertEqual(rc, px.EXIT_USAGE)
            self.assertEqual(list(Path(td).iterdir()), [], "nothing may be written when a format is unusable")


class TestPresentationSwitches(unittest.TestCase):
    """Presentation is adjustable; safety is not."""

    def _run(self, **kw):
        td = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, td, ignore_errors=True)
        cfg = make_cfg(output=Path(td) / "out", **kw)
        with run_with(FakeClient()):
            px.run(cfg)
        return Path(td)

    def _csv(self, td, name="out.csv"):
        with (td / name).open(encoding="utf-8-sig") as fh:
            return list(csv.reader(fh))

    # --- headers ---------------------------------------------------------------------
    def test_readable_rule(self):
        self.assertEqual(px.readable_header("archive_serial_number"), "Archive serial number")
        self.assertEqual(px.readable_header("file:version"), "File version")
        self.assertEqual(px.readable_header("custom_field:Amount"), "Amount")
        self.assertEqual(px.readable_header("custom_field:Amount:currency"), "Amount currency")
        self.assertEqual(px.readable_header("custom_field:Invoice Number"), "Invoice Number")
        self.assertEqual(px.readable_header("custom_field:Amount#7"), "Amount#7")

    def test_colliding_readable_headings_fall_back_to_raw(self):
        cols = ["id", "notes", "custom_field:Notes", "custom_field:Amount"]
        self.assertEqual(px.readable_clashes(cols), ["Notes"])
        self.assertEqual(px.headings(cols, "readable"),
                         ["Id", "notes", "custom_field:Notes", "Amount"])

    def test_readable_is_the_csv_default_and_raw_is_the_json_default(self):
        td = self._run(formats=["csv", "json"])
        self.assertIn("Archive serial number", self._csv(td)[0])
        payload = json.loads((td / "out.json").read_text(encoding="utf-8"))
        self.assertIn("archive_serial_number", payload["rows"][0])

    def test_raw_switch_changes_the_csv_header(self):
        td = self._run(formats=["csv"], headers="raw")
        head = self._csv(td)[0]
        self.assertIn("archive_serial_number", head)
        self.assertIn("custom_field:Amount", head)

    def test_readable_switch_changes_the_json_keys(self):
        td = self._run(formats=["json", "jsonl"], headers="readable")
        payload = json.loads((td / "out.json").read_text(encoding="utf-8"))
        self.assertIn("Archive serial number", payload["rows"][0])
        first = json.loads((td / "out.jsonl").read_text(encoding="utf-8").splitlines()[0])
        self.assertIn("Amount", first)

    def test_columns_still_take_raw_names_under_readable_headers(self):
        td = self._run(formats=["csv"], columns=["id", "title", "custom_field:Amount"])
        self.assertEqual(self._csv(td)[0], ["Id", "Title", "Amount"])

    def test_trailer_records_the_header_style_per_file(self):
        td = self._run(formats=["csv", "json"])
        self.assertIn("# headers: readable", (td / "out.csv").read_text(encoding="utf-8-sig"))
        payload = json.loads((td / "out.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["trailer"]["headers"], "raw")

    # --- date ------------------------------------------------------------------------
    def test_date_presets(self):
        self.assertEqual(px.fmt_date("2025-01-05T00:00:00Z", px.resolve_date_format("iso")), "2025-01-05")
        self.assertEqual(px.fmt_date("2025-01-05T00:00:00Z", px.resolve_date_format("us")), "01/05/2025")
        self.assertEqual(px.fmt_date("2025-01-05T00:00:00Z", px.resolve_date_format("eu")), "05.01.2025")

    def test_custom_strftime_pattern_is_accepted(self):
        fmt = px.resolve_date_format("%d %b %Y")
        self.assertEqual(px.fmt_date("2025-01-05T00:00:00Z", fmt), "05 Jan 2025")

    def test_date_applies_in_every_format_and_to_date_custom_fields(self):
        td = self._run(formats=["csv", "json"], date_spec="us", date_fmt=px.resolve_date_format("us"))
        rows = self._csv(td)
        created = rows[1][rows[0].index("Created")]
        self.assertEqual(created, "01/05/2025")
        payload = json.loads((td / "out.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["rows"][0]["created"], "01/05/2025")

    def test_date_pattern_that_makes_formula_bait_is_refused(self):
        with self.assertRaises(px.UsageError):
            px.resolve_date_format("=%Y")
        with self.assertRaises(px.UsageError):
            px.resolve_date_format("-%Y-%m-%d")

    def test_pattern_that_substitutes_nothing_is_refused(self):
        for bad in ("%Q", "invoices"):
            with self.assertRaises(px.UsageError, msg=bad):
                px.resolve_date_format(bad)

    def test_empty_date_pattern_is_refused(self):
        with self.assertRaises(px.UsageError):
            px.resolve_date_format("")

    def test_trailer_records_the_date_choice(self):
        td = self._run(formats=["csv"], date_spec="us", date_fmt=px.resolve_date_format("us"))
        self.assertIn("# date format: us", (td / "out.csv").read_text(encoding="utf-8-sig"))

    # --- list separator --------------------------------------------------------------
    def test_default_list_separator(self):
        td = self._run(formats=["csv"])
        rows = self._csv(td)
        self.assertEqual(rows[1][rows[0].index("Tags")], "tax; 2025")

    def test_list_separator_applies_to_tags_and_documentlink(self):
        td = self._run(formats=["csv"], list_sep=" | ")
        rows = self._csv(td)
        self.assertEqual(rows[1][rows[0].index("Tags")], "tax | 2025")

    def test_list_separator_equal_to_the_csv_separator_is_refused(self):
        import os
        old = {k: os.environ.get(k) for k in ("PAPERLESS_URL", "PAPERLESS_TOKEN")}
        os.environ["PAPERLESS_URL"] = "https://paperless.test"
        os.environ["PAPERLESS_TOKEN"] = "t"
        try:
            with self.assertRaises(px.UsageError):
                px.config_from_args(["--list-sep", ","])
            with self.assertRaises(px.UsageError):
                px.config_from_args(["--locale", "comma", "--list-sep", ";"])
            px.config_from_args(["--locale", "comma", "--list-sep", ","])  # legal: not that locale's separator
        finally:
            for k, v in old.items():
                os.environ.pop(k, None) if v is None else os.environ.__setitem__(k, v)

    # --- what no switch may touch ----------------------------------------------------
    def test_switches_cannot_disable_the_formula_guard(self):
        td = self._run(formats=["csv"], headers="raw", list_sep=" | ",
                       date_spec="us", date_fmt=px.resolve_date_format("us"))
        rows = self._csv(td)
        titles = [r[rows[0].index("title")] for r in rows[1:] if len(r) > 1]
        self.assertIn("'=1+1", titles)

    def test_switches_cannot_disable_the_row_count_check(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = make_cfg(output=Path(td) / "out", headers="raw", list_sep=" | ")
            with run_with(FakeClient(count=99)):
                with self.assertRaises(px.ReconcileError):
                    px.run(cfg)


class TestZeroMatches(unittest.TestCase):
    def test_empty_result_is_success_with_headers(self):
        with tempfile.TemporaryDirectory() as td:
            cfg = make_cfg(output=Path(td) / "out", formats=["csv"])
            with run_with(FakeClient(docs=[], count=0)):
                rc = px.run(cfg)
            self.assertEqual(rc, 0)
            with (Path(td) / "out.csv").open(encoding="utf-8-sig") as fh:
                rows = list(csv.reader(fh))
            self.assertEqual(rows[0][0], "Id")


class TestVersion(unittest.TestCase):
    """One version number, pinned everywhere it can be checked. A release with two answers is a
    release we cannot support.

    pyproject.toml and CHANGELOG.md both ship flat beside this file's module in the public repo,
    but sit one level up in this project's own layout — gen_rules.find_ship_doc resolves either.
    """

    here = Path(__file__).resolve().parent

    def test_version_flag_prints_the_version(self):
        out = io.StringIO()
        with contextlib.redirect_stdout(out), self.assertRaises(SystemExit) as cm:
            px.build_parser().parse_args(["--version"])
        self.assertEqual(cm.exception.code, 0)
        self.assertEqual(out.getvalue().strip(), f"paperless-export {px.__version__}")

    def test_pyproject_agrees(self):
        path = gen_rules.find_ship_doc("pyproject.toml", self.here)
        self.assertIsNotNone(path, "pyproject.toml not found beside or above this file")
        text = path.read_text(encoding="utf-8")
        m = re.search(r'^version\s*=\s*"([^"]+)"', text, re.M)
        self.assertIsNotNone(m, "pyproject.toml has no version")
        self.assertEqual(m.group(1), px.__version__)

    def test_changelog_top_entry_agrees(self):
        path = gen_rules.find_ship_doc("CHANGELOG.md", self.here)
        self.assertIsNotNone(path, "CHANGELOG.md not found beside or above this file")
        text = path.read_text(encoding="utf-8")
        m = re.search(r"^##\s*\[([0-9]+\.[0-9]+\.[0-9]+)\]", text, re.M)
        self.assertIsNotNone(m, "CHANGELOG.md has no [X.Y.Z] version heading")
        self.assertEqual(m.group(1), px.__version__)


class TestShipSetIsSterile(unittest.TestCase):
    """Nothing that ships names a private detail of how this tool was built: not a person, not
    who decided something, not this project's own private files, not this machine's path. The
    public changelog says what changed; it never says who decided it or where the discussion
    happened. One pattern, run over every file that actually ships, naming the file and line when
    it is not empty -- the same discipline as the rule-map check above, extended to prose."""

    here = Path(__file__).resolve().parent

    # This line necessarily spells out what it forbids, so the scan below excludes it by an
    # exact marker rather than by accident of wording -- everything else in this file is fair
    # game, including this class's own comments and docstrings.
    STERILITY_PATTERN = re.compile(r"jody|bunky|spaghettio|\bGM\b|\bcoder\b|\bset [0-9]|ruling|decisions\.md|progress\.md|memory/|architecture/|CLAUDE|\bSOP\b|§|\bexport\.md\b|\.tmp|/Users/|Lab/|claude|anthropic|co-authored|generated with", re.IGNORECASE)  # STERILITY_PATTERN_DEFINITION

    SHIPPED_MODULES = ("paperless_export.py", "paperless_serve.py", "paperless_front.py",
                        "gen_rules.py", "probe_paperless.py", "fixture_v10.json",
                        "test_paperless_export.py", "test_front2.py", "test_front2_batter.py",
                        "test_front2_conformance.py", "test_front3.py", "test_front3_batter.py")
    SHIPPED_DOCS = ("README.md", "LICENSE", "NOTICE", "SECURITY.md", "CHANGELOG.md",
                     "pyproject.toml", "requirements.txt")

    def _ship_files(self) -> list[Path]:
        files = [self.here / name for name in self.SHIPPED_MODULES]
        for name in self.SHIPPED_DOCS:
            path = gen_rules.find_ship_doc(name, self.here)
            if path is not None:
                files.append(path)
        # GITIGNORE is this project's own staging name for the file that becomes .gitignore on
        # transfer -- a different, real .gitignore already exists in this project's own layout,
        # governing files this repository never ships. Checked in that order so the private one
        # is never mistaken for the public one.
        gitignore = (gen_rules.find_ship_doc("GITIGNORE", self.here)
                     or gen_rules.find_ship_doc(".gitignore", self.here))
        if gitignore is not None:
            files.append(gitignore)
        return [f for f in files if f.is_file()]

    def test_no_shipped_file_names_how_this_tool_was_built(self):
        hits = []
        for path in self._ship_files():
            for lineno, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
                if "STERILITY_PATTERN_DEFINITION" in line:
                    continue
                if self.STERILITY_PATTERN.search(line):
                    hits.append(f"{path.name}:{lineno}: {line.strip()}")
        self.assertEqual(hits, [],
                          "shipped file names a private detail of how this tool was built:\n"
                          + "\n".join(hits))


class TestTestsNameTheirEncoding(unittest.TestCase):
    """A read_text/write_text/open/subprocess call with no explicit encoding decodes with the OS
    locale codepage on Windows, not UTF-8 -- found and fixed by hand more than once. One scan over
    every test file's own source closes the class of bug instead of the instances of it that were
    actually seen.
    """

    # Matches the call name up to its opening paren; the paren body is walked separately below
    # because it can nest (a path, dot, the read call; subprocess, dot, run with a list, etc).
    CALL = re.compile(
        r"(?<![\w.])open\(|\.read_text\(|\.write_text\("
        r"|subprocess\.run\(|subprocess\.Popen\(|subprocess\.check_output\("
    )
    BINARY_MODE = re.compile(r"""["'][^"']*[abrwx][b]["']""")

    @staticmethod
    def _code_only(src: str) -> str:
        """Blank out comments so a sentence describing a call (this class's own comments,
        included) cannot be mistaken for the call itself. Positions are preserved -- only comment
        characters become spaces -- so a real match's reported line number stays accurate.
        """
        chars = list(src)
        line_start = [0]
        for line in src.splitlines(keepends=True):
            line_start.append(line_start[-1] + len(line))
        try:
            for tok in tokenize.generate_tokens(io.StringIO(src).readline):
                if tok.type != tokenize.COMMENT:
                    continue
                (srow, scol), (erow, ecol) = tok.start, tok.end
                for i in range(line_start[srow - 1] + scol, line_start[erow - 1] + ecol):
                    chars[i] = " "
        except tokenize.TokenizeError:
            pass
        return "".join(chars)

    def _call_body(self, src: str, open_paren: int) -> str:
        depth, i = 1, open_paren
        while depth and i < len(src) - 1:
            i += 1
            if src[i] == "(":
                depth += 1
            elif src[i] == ")":
                depth -= 1
        return src[open_paren + 1:i]

    def _violations(self, path: Path) -> list[str]:
        src = path.read_text(encoding="utf-8")
        scan = self._code_only(src)
        out = []
        for m in self.CALL.finditer(scan):
            body = self._call_body(scan, m.end() - 1)
            if "encoding=" in body:
                continue
            is_open = m.group(0) == "open(" or m.group(0) == ".open("
            if is_open and self.BINARY_MODE.search(body):
                continue  # binary mode takes no encoding
            is_subprocess = m.group(0).startswith("subprocess.")
            if is_subprocess and "text=True" not in body and "universal_newlines=True" not in body:
                continue  # bytes in, bytes out -- no implicit decode to drift by locale
            line = src.count("\n", 0, m.start()) + 1
            out.append(f"{path.name}:{line}: {m.group(0)}...) with no encoding=")
        return out

    def test_every_test_file_names_its_encoding(self):
        here = Path(__file__).resolve().parent
        problems = []
        for f in sorted(here.glob("test_*.py")):
            problems.extend(self._violations(f))
        self.assertEqual(problems, [], "\n" + "\n".join(problems))


class TestPrintedOutputIsASCII(unittest.TestCase):
    """Windows decodes a child process's console output with its own codepage, not UTF-8 -- an
    em dash written into the --serve banner crashed a caller that read the pipe as UTF-8. A
    document title, a saved-view name, or anything else read from the user's own paperless
    instance is not ours to restrict and is excluded here; only the literal text this project
    typed into a print(...) or sys.stderr.write(...) call is checked, never an f-string's
    placeholder expressions, which carry that runtime data.
    """

    MODULES = ("paperless_export.py", "paperless_serve.py", "paperless_front.py",
               "gen_rules.py", "probe_paperless.py")

    @staticmethod
    def _is_target_call(node: ast.Call) -> bool:
        f = node.func
        if isinstance(f, ast.Name) and f.id == "print":
            return True
        return (isinstance(f, ast.Attribute) and f.attr == "write"
                and isinstance(f.value, ast.Attribute) and f.value.attr == "stderr")

    @staticmethod
    def _literal_text(node: ast.expr):
        """The hand-written text of one printed argument: a plain string as itself, or an
        f-string's own Constant pieces -- never a FormattedValue, which is a runtime substitution."""
        if isinstance(node, ast.Constant) and isinstance(node.value, str):
            yield node.value
        elif isinstance(node, ast.JoinedStr):
            for piece in node.values:
                if isinstance(piece, ast.Constant) and isinstance(piece.value, str):
                    yield piece.value

    def _violations(self, path: Path) -> list[str]:
        tree = ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
        out = []
        for node in ast.walk(tree):
            if not (isinstance(node, ast.Call) and self._is_target_call(node)):
                continue
            for arg in node.args:
                for text in self._literal_text(arg):
                    if any(ord(c) > 127 for c in text):
                        out.append(f"{path.name}:{node.lineno}: non-ASCII character in {text!r}")
        return out

    def test_no_shipped_module_prints_a_non_ascii_character(self):
        here = Path(__file__).resolve().parent
        problems = []
        for name in self.MODULES:
            path = here / name
            if path.is_file():
                problems.extend(self._violations(path))
        self.assertEqual(problems, [], "\n" + "\n".join(problems))


if __name__ == "__main__":
    unittest.main(verbosity=2)
